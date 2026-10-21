"""Task 2 - Excel/CSV bulk importer (SRS REQ-1.1, 1.2, 1.4).

Rules:
  * Pehle POORI file validate hoti hai. Ek bhi error ho to DB me KUCH save nahi hota (REQ-1.2).
  * Errors row + column ke saath milte hain.
  * Same file dobara upload = update (duplicate nahi banta).
Upload order: term -> department -> timeslot -> room -> faculty -> course -> student -> choice -> preference
"""
import csv
import datetime as dt
import io
import re
from dataclasses import dataclass, field
from pathlib import Path

from django.db import transaction

from . import validators as V
from .models import (AcademicTerm, Course, Department, Faculty, FacultyPreference, Room, Student,
                     StudentCourseChoice, TimeSlot)


@dataclass
class Col:
    name: str
    parser: object
    required: bool = True
    default: object = None
    target: str = None          # model field ka naam (default = column name)
    example: str = ""
    help: str = ""


@dataclass
class Spec:
    model: type
    key: tuple                  # natural key (model field names)
    columns: list
    cross: object = None        # cross-field validator
    description: str = ""


SPECS = {
    "term": Spec(AcademicTerm, ("name",), [
        Col("name", V.s(20), example="2026-ODD"),
        Col("is_active", V.b(), False, False, example="yes", help="yes = active term (choice import / solver ka default)")],
        description="Academic term / semester"),
    "department": Spec(Department, ("code",), [
        Col("code", V.s(10), example="CSE"), Col("name", V.s(100), example="Computer Science")],
        description="Departments"),
    "timeslot": Spec(TimeSlot, ("day", "period"), [
        Col("day", V.choice(V.DAYS), example="Mon", help="Mon..Sat"),
        Col("period", V.i(1, 20), example="1"),
        Col("start_time", V.t(), example="09:00"), Col("end_time", V.t(), example="10:00"),
        Col("is_break", V.b(), False, False, example="no", help="yes = lunch/break, yahan class nahi lagegi")],
        V.check_timeslot, "Working periods (aur breaks)"),
    "room": Spec(Room, ("code",), [
        Col("code", V.s(20), example="R101"), Col("name", V.s(100), False, "", example="Room 101"),
        Col("capacity", V.i(1), example="60"),
        Col("room_type", V.choice(V.ROOM_TYPES), False, "CLASSROOM", example="classroom", help="classroom / lab / seminar")],
        description="Rooms aur labs"),
    "faculty": Spec(Faculty, ("employee_id",), [
        Col("employee_id", V.s(20), example="F001"), Col("name", V.s(100), example="Prof. Sharma"),
        Col("email", V.email(), False, "", example="sharma@college.edu"),
        Col("department", V.fk(Department, "code", "Department"), target="department", example="CSE", help="Department code"),
        Col("max_hours_per_week", V.i(1, 60), False, 16, example="16")],
        description="Faculty"),
    "course": Spec(Course, ("code",), [
        Col("code", V.s(20), example="CS301"), Col("name", V.s(150), example="DBMS"),
        Col("department", V.fk(Department, "code", "Department"), example="CSE", help="Department code"),
        Col("course_type", V.choice(V.COURSE_TYPES), False, "MAJOR", example="major", help="major / minor / skill / value-added"),
        Col("credits", V.i(0, 40), False, 3, example="3"),
        Col("hours_per_week", V.i(1, 40), False, 3, example="3"),
        Col("requires_lab", V.b(), False, False, example="no"),
        Col("lab_block_size", V.i(1, 8), False, 1, example="1", help="Lab ke consecutive periods (1 = normal)"),
        Col("semester", V.i(1, 12), False, 1, example="5"),
        Col("faculty_id", V.fk(Faculty, "employee_id", "Faculty"), False, None, target="faculty", example="F001",
            help="Faculty ka employee_id (khali = baad me assign)"),
        Col("expected_students", V.i(0), False, 0, example="55")],
        V.check_course, "Courses (NEP: major/minor/skill/value-added)"),
    "student": Spec(Student, ("roll_no",), [
        Col("roll_no", V.s(20), example="2026CS001"), Col("name", V.s(100), example="Asha Verma"),
        Col("email", V.email(), False, "", example="asha@college.edu"),
        Col("department", V.fk(Department, "code", "Department"), example="CSE"),
        Col("semester", V.i(1, 12), False, 1, example="5")],
        description="Students"),
    "choice": Spec(StudentCourseChoice, ("student", "course", "term"), [
        Col("student_roll_no", V.fk(Student, "roll_no", "Student"), target="student", example="2026CS001"),
        Col("course_code", V.fk(Course, "code", "Course"), target="course", example="CS301"),
        Col("term", V.fk(AcademicTerm, "name", "Term"), False, None, example="2026-ODD",
            help="Khali = active term (ya --term)")],
        description="Student ke elective / minor / skill / value-added choices"),
    "preference": Spec(FacultyPreference, ("faculty", "timeslot"), [
        Col("faculty_id", V.fk(Faculty, "employee_id", "Faculty"), target="faculty", example="F001"),
        Col("day", V.choice(V.DAYS), example="Mon"), Col("period", V.i(1, 20), example="1"),
        Col("level", V.choice(V.LEVELS), example="NO", help="NO (unavailable) / AVOID / PREFER")],
        description="Faculty availability aur preferences"),
}
ORDER = ["term", "department", "timeslot", "room", "faculty", "course", "student", "choice", "preference"]


# ------------------------------------------------------------------ file reading
def _norm_header(h):
    return re.sub(r"[\s\-]+", "_", str(h).strip().lower())


def _cell(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, dt.datetime):
        return v.strftime("%H:%M") if v.year <= 1900 else v.strftime("%Y-%m-%d")
    if isinstance(v, dt.time):
        return v.strftime("%H:%M")
    return str(v).strip()


def read_table(file, filename=None):
    """Returns (headers, rows[(sheet_row_no, {header: str})]). file = path ya file-like."""
    if hasattr(file, "read"):
        name = filename or getattr(file, "name", "")
        raw = file.read()
    else:
        name, raw = str(file), Path(file).read_bytes()
    ext = Path(name).suffix.lower()
    if ext == ".csv":
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("latin-1")
        first = text.splitlines()[0] if text.strip() else ""
        delim = ";" if first.count(";") > first.count(",") else ","
        grid = list(csv.reader(io.StringIO(text), delimiter=delim))
    elif ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        ws = wb["data"] if "data" in wb.sheetnames else wb.worksheets[0]
        grid = [list(r) for r in ws.iter_rows(values_only=True)]
    else:
        raise ValueError("Sirf .xlsx ya .csv file upload karo")
    if not grid:
        raise ValueError("File khali hai")
    headers = [_norm_header(h) for h in grid[0]]
    rows = []
    for n, r in enumerate(grid[1:], start=2):
        cells = [_cell(c) for c in r]
        if any(cells):
            rows.append((n, {h: (cells[k] if k < len(cells) else "") for k, h in enumerate(headers) if h}))
    return headers, rows


# ------------------------------------------------------------------ importer
class Ctx:
    def __init__(self, default_term=None):
        self.default_term, self._cache = default_term, {}

    def lookup(self, model, field):
        key = (model, field)
        if key not in self._cache:
            self._cache[key] = {str(getattr(o, field)): o for o in model.objects.all()}
        return self._cache[key]


@dataclass
class ImportResult:
    entity: str
    total_rows: int = 0
    created: int = 0
    updated: int = 0
    errors: list = field(default_factory=list)      # [{"row", "column", "message"}]
    warnings: list = field(default_factory=list)
    dry_run: bool = False

    @property
    def ok(self):
        return not self.errors

    def summary(self):
        if not self.ok:
            return f"{self.entity}: {len(self.errors)} error(s) in {self.total_rows} rows - KUCH save nahi hua"
        verb = "validated (dry run)" if self.dry_run else f"created {self.created}, updated {self.updated}"
        return f"{self.entity}: {self.total_rows} rows OK - {verb}"

    def as_dict(self):
        return {"entity": self.entity, "ok": self.ok, "total_rows": self.total_rows, "created": self.created,
                "updated": self.updated, "errors": self.errors, "warnings": self.warnings, "dry_run": self.dry_run}


def import_file(entity, file, *, term=None, dry_run=False, filename=None):
    """term: AcademicTerm ya None (choice import ke liye default). Returns ImportResult."""
    if entity not in SPECS:
        raise ValueError(f"Unknown entity '{entity}'. Allowed: {', '.join(ORDER)}")
    spec, res = SPECS[entity], ImportResult(entity, dry_run=dry_run)
    try:
        headers, rows = read_table(file, filename)
    except ValueError as e:
        res.errors.append({"row": 1, "column": "", "message": str(e)})
        return res
    res.total_rows = len(rows)
    missing = [c.name for c in spec.columns if c.required and c.name not in headers]
    if missing:
        res.errors.append({"row": 1, "column": ", ".join(missing), "message": "Required column(s) header me nahi mile"})
        return res
    known = {c.name for c in spec.columns}
    extra = [h for h in headers if h and h not in known]
    if extra:
        res.warnings.append(f"Ignored unknown column(s): {', '.join(extra)}")
    if not rows:
        res.errors.append({"row": 2, "column": "", "message": "Koi data row nahi mili"})
        return res

    ctx = Ctx()
    ctx.default_term = term or AcademicTerm.objects.filter(is_active=True).order_by("-id").first()
    cleaned_rows, seen = [], {}
    slots = {(x.day, x.period): x for x in TimeSlot.objects.all()} if entity == "preference" else {}
    for n, raw in rows:
        clean, row_ok = {}, True
        for col in spec.columns:
            if col.name not in headers:      # optional column sheet me hai hi nahi -> existing value ko mat chhedo
                continue
            v = raw.get(col.name, "")
            if v == "":
                if col.required:
                    res.errors.append({"row": n, "column": col.name, "message": "Required hai, khali nahi ho sakta"})
                    row_ok = False
                    continue
                clean[col.target or col.name] = col.default
                continue
            try:
                clean[col.target or col.name] = col.parser(v, ctx)
            except ValueError as e:
                res.errors.append({"row": n, "column": col.name, "message": str(e)})
                row_ok = False
        if not row_ok:
            continue
        if entity == "choice" and clean.get("term") is None:
            if ctx.default_term is None:
                res.errors.append({"row": n, "column": "term", "message": "term khali hai aur koi active term nahi hai"})
                continue
            clean["term"] = ctx.default_term
        if entity == "preference":
            ts = slots.get((clean.pop("day"), clean.pop("period")))
            if ts is None:
                res.errors.append({"row": n, "column": "day/period", "message": "Is day+period ka TimeSlot nahi hai (pehle timeslot import karo)"})
                continue
            clean["timeslot"] = ts
        if spec.cross:
            bad = spec.cross(clean)
            for column, msg in bad:
                res.errors.append({"row": n, "column": column, "message": msg})
            if bad:
                continue
        key = tuple(getattr(clean[k], "pk", clean[k]) for k in spec.key)
        if key in seen:
            res.errors.append({"row": n, "column": ", ".join(spec.key),
                               "message": f"Duplicate - yahi record row {seen[key]} me bhi hai"})
            continue
        seen[key] = n
        cleaned_rows.append(clean)

    if res.errors or dry_run:
        return res
    with transaction.atomic():
        for clean in cleaned_rows:
            lookup = {k: clean[k] for k in spec.key}
            defaults = {k: v for k, v in clean.items() if k not in spec.key}
            _, created = spec.model.objects.update_or_create(defaults=defaults, **lookup)
            res.created += created
            res.updated += (not created)
        if entity == "choice":
            _refresh_expected_students({c["term"] for c in cleaned_rows})
    return res


def _refresh_expected_students(terms):
    """Course.expected_students = max(purana, is term ke choices ki ginti). Room capacity check isi se hota hai."""
    from django.db.models import Count
    for term in terms:
        for row in StudentCourseChoice.objects.filter(term=term).values("course_id").annotate(n=Count("id")):
            Course.objects.filter(pk=row["course_id"], expected_students__lt=row["n"]).update(expected_students=row["n"])


# ------------------------------------------------------------------ templates
def build_template(entity):
    """Blank .xlsx template (header + 1 example row + Help sheet) bytes me."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    spec = SPECS[entity]
    wb = Workbook()
    ws = wb.active
    ws.title = "data"
    ws.append([c.name for c in spec.columns])
    ws.append([c.example for c in spec.columns])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for idx, c in enumerate(spec.columns, start=1):
        ws.column_dimensions[ws.cell(1, idx).column_letter].width = max(14, len(c.name) + 4)
    ws.append([])
    helps = wb.create_sheet("help")
    helps.append(["column", "required", "default", "notes"])
    for c in spec.columns:
        helps.append([c.name, "yes" if c.required else "no", "" if c.default in (None, "") else str(c.default), c.help])
    helps.append([])
    helps.append([f"Example row delete karke apna data likho. Sheet ka naam 'data' rakho. Upload order: {' -> '.join(ORDER)}"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()